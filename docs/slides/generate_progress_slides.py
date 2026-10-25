"""Generate the progress-review deck: Bobas_Vasylchenko_Hombosh_Progress.pptx / .pdf.

Run from the repository root:

    python docs/slides/generate_progress_slides.py

Needs ``python-pptx``; the PDF is exported with LibreOffice (``soffice``) if it is
installed, otherwise open the .pptx in PowerPoint / Google Slides and export to PDF.
Every number on the slides comes from ``results/experiments.md`` (notebook runs of
25.10.2026). Speaker notes hold the talk track and who presents which slide.
"""

import shutil
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
FIGURES = HERE / "figures"
OUT_PPTX = HERE / "Bobas_Vasylchenko_Hombosh_Progress.pptx"
TOTAL = 7
FONT = "Calibri"

BG = RGBColor(0x0B, 0x0F, 0x19)
CARD = RGBColor(0x15, 0x1C, 0x2E)
CARD_LINE = RGBColor(0x23, 0x2F, 0x48)
WHITE = RGBColor(0xF1, 0xF5, 0xF9)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
SKY = RGBColor(0x38, 0xBD, 0xF8)
GREEN = RGBColor(0x10, 0xB9, 0x81)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
ROSE = RGBColor(0xF4, 0x3F, 0x5E)
VIOLET = RGBColor(0xA7, 0x8B, 0xFA)
DARK_TEXT = RGBColor(0x0B, 0x0F, 0x19)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# --------------------------------------------------------------------------- helpers


def _style_run(run, size, color, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic


def _bullet(paragraph, char="•", color=SKY):
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set("marL", str(Inches(0.22)))
    p_pr.set("indent", str(-Inches(0.22)))
    bu_clr = p_pr.makeelement(qn("a:buClr"), {})
    srgb = bu_clr.makeelement(qn("a:srgbClr"), {"val": str(color)})
    bu_clr.append(srgb)
    p_pr.append(bu_clr)
    p_pr.append(p_pr.makeelement(qn("a:buChar"), {"char": char}))


def text(
    slide,
    x,
    y,
    w,
    h,
    paragraphs,
    size=14,
    color=WHITE,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    bullets=False,
    space_after=6,
):
    """paragraphs: list of str or list of [(text, {bold, color, size, italic}), ...]."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = [(para, {})] if isinstance(para, str) else para
        for chunk, opt in runs:
            r = p.add_run()
            r.text = chunk
            _style_run(
                r,
                opt.get("size", size),
                opt.get("color", color),
                opt.get("bold", False),
                opt.get("italic", False),
            )
        if bullets:
            _bullet(p)
    return box


def card(slide, x, y, w, h, fill=CARD, line=CARD_LINE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


def pill(slide, x, y, label, fill, w=None, size=10, color=DARK_TEXT):
    w = w or 0.085 * len(label) + 0.45
    shape = card(slide, x, y, w, 0.3, fill=fill, line=None)
    shape.adjustments[0] = 0.5
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    _style_run(r, size, color, bold=True)
    return shape


def picture(slide, path, x, y, w=None, h=None):
    kwargs = {}
    if w:
        kwargs["width"] = Inches(w)
    if h:
        kwargs["height"] = Inches(h)
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kwargs)


def _cell_borders(cell, color=CARD_LINE, width=Pt(1)):
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnB", "a:lnT", "a:lnR", "a:lnL"):
        ln = tc_pr.makeelement(qn(tag), {"w": str(int(width)), "cmpd": "sng"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        fill.append(fill.makeelement(qn("a:srgbClr"), {"val": str(color)}))
        ln.append(fill)
        tc_pr.insert(0, ln)


def table(
    slide, x, y, w, rows, col_widths, row_h=0.42, size=13, highlight_row=None, left_cols=(0,)
):
    shape = slide.shapes.add_table(
        len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows))
    )
    tbl = shape.table
    for j, cw in enumerate(col_widths):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = CARD_LINE
            elif i == highlight_row:
                cell.fill.fore_color.rgb = RGBColor(0x0F, 0x3A, 0x33)
            else:
                cell.fill.fore_color.rgb = CARD
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j in left_cols else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = value
            _style_run(r, size, MUTED if i == 0 else WHITE, bold=(i == 0 or i == highlight_row))
            _cell_borders(cell)
    tbl_pr = shape._element.graphic.graphicData.tbl.tblPr
    tbl_pr.set("bandRow", "0")
    tbl_pr.set("firstRow", "0")
    return shape


def arrow(slide, x, y, w=0.45, h=0.32, color=SKY):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def new_slide(num, badge, title, subtitle, badge_color=SKY, notes=""):
    slide = prs.slides.add_slide(BLANK)
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    pill(slide, 0.6, 0.4, badge, badge_color)
    text(slide, 0.6, 0.82, 12.1, 0.6, [[(title, {"bold": True})]], size=30)
    if subtitle:
        text(slide, 0.6, 1.45, 12.1, 0.4, [subtitle], size=16, color=MUTED)
    text(
        slide,
        0.6,
        7.0,
        8,
        0.3,
        ["SkyNet · Бобас, Васильченко, Гомбош · AI course, УКУ 2026"],
        size=10,
        color=MUTED,
    )
    text(
        slide, 10.73, 7.0, 2, 0.3, [f"{num} / {TOTAL}"], size=10, color=MUTED, align=PP_ALIGN.RIGHT
    )
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


# --------------------------------------------------------------------------- 1. title

s = new_slide(
    1,
    "AI COURSE · ПРОМІЖНИЙ ПРОГРЕС · 2026",
    "Яке представлення сцени найкраще для world model?",
    "Порівнюємо object-centric, BEV і latent BEV на nuScenes за однаковими правилами оцінки",
    notes=(
        "Доповідає: Юрій. 30 с. World model прогнозує, як зміниться дорожня сцена в наступні "
        "секунди. Ми порівнюємо три способи описати сцену: як список об'єктів, як карту зверху "
        "(BEV) і як стиснену карту (latent BEV), і перевіряємо, з яким прогноз на 1–3 секунди "
        "найточніший. Перше представлення готове, друге зібране й навчене, третє — наступний етап."
    ),
)
card(s, 0.6, 2.2, 5.6, 4.0)
text(
    s,
    0.9,
    2.5,
    5.1,
    4.1,
    [
        [("Команда", {"bold": True, "color": SKY, "size": 17})],
        "Юрій Бобас · Владислав Васильченко",
        "Олег Гомбош",
        [("Менторка", {"bold": True, "color": SKY, "size": 17})],
        "Марта Сумик",
        [("Репозиторій", {"bold": True, "color": SKY, "size": 17})],
        "github.com/Hombosh-dev/SkyNet",
        [("Дані", {"bold": True, "color": SKY, "size": 17})],
        "nuScenes mini: 10 сцен по 20 с, 6 камер + LiDAR",
    ],
    size=16,
    space_after=7,
)
picture(s, FIGURES / "surround_cameras.jpg", 6.6, 2.2, w=6.13)
text(
    s,
    6.6,
    2.2 + 6.13 * 540 / 1440 + 0.08,
    6.13,
    0.3,
    ["6 камер одного моменту з nuScenes"],
    size=13,
    color=MUTED,
)
y_status = 2.2 + 6.13 * 540 / 1440 + 0.75
for k, (name, label, color) in enumerate(
    [
        ("Object-centric", "ГОТОВО", GREEN),
        ("BEV", "ДОНАВЧАННЯ", AMBER),
        ("Latent BEV", "ПЛАНУЄТЬСЯ", ROSE),
    ]
):
    x = 6.6 + k * 2.1
    card(s, x, y_status, 1.95, 0.95)
    text(s, x + 0.15, y_status + 0.1, 1.7, 0.35, [[(name, {"bold": True})]], size=16)
    pill(s, x + 0.15, y_status + 0.52, label, color, w=1.4, size=11)

# --------------------------------------------------------------------------- 2. pipeline

s = new_slide(
    2,
    "ТЕХНІЧНА ЧАСТИНА · PIPELINE",
    "Pipeline і чому саме ці моделі",
    "Усі представлення отримують ті самі кадри й оцінюються тими самими метриками",
    notes=(
        "Доповідає: Владислав. 45 с. Беремо nuScenes mini: 8 сцен для навчання, 2 для тесту. "
        "Модель бачить 3 минулі кадри і прогнозує 6 наступних, тобто 3 секунди. Object-centric "
        "описує кожного учасника руху кількома числами — це найкомпактніше і найзрозуміліше. BEV — "
        "карта зверху, на яку зводимо 6 камер і лідар; ми взяли архітектуру BEVFusion, бо це "
        "поширений підхід і він вміщається в GPU Colab. Latent BEV — та сама карта, стиснена "
        "автоенкодером у латентний простір, де прогнозувати дешевше."
    ),
)
flow = [
    ("nuScenes mini", "8 сцен — навчання\n2 сцени — тест"),
    ("3 минулі кадри", "камери, LiDAR,\nанотації, рух авто"),
    ("Представлення", "object-centric · BEV ·\nlatent BEV"),
    ("Прогноз", "6 кадрів = 3 с"),
    ("Оцінка", "однакові метрики\nі бейзлайни"),
]
bw, gap, y0 = 2.05, 0.46, 2.15
for k, (head, body) in enumerate(flow):
    x = 0.6 + k * (bw + gap)
    card(s, x, y0, bw, 1.35, fill=CARD if k != 2 else RGBColor(0x12, 0x2A, 0x45))
    text(s, x + 0.15, y0 + 0.15, bw - 0.3, 0.35, [[(head, {"bold": True, "color": SKY})]], size=16)
    text(s, x + 0.15, y0 + 0.58, bw - 0.3, 0.75, body.split("\n"), size=14, space_after=0)
    if k < len(flow) - 1:
        arrow(s, x + bw + 0.07, y0 + 0.52)

models = [
    (
        "Object-centric",
        GREEN,
        "кожен агент — вектор: положення, розмір, курс, швидкість, клас",
        "найкомпактніший та інтерпретований опис сцени",
    ),
    (
        "BEV (BEVFusion)",
        AMBER,
        "6 камер (Lift-Splat) + LiDAR (PointPillars) в одну карту зверху 100\u00a0×\u00a0100\u00a0м",
        "поширений підхід до злиття сенсорів; вміщається в GPU Colab",
    ),
    (
        "Latent BEV",
        ROSE,
        "та сама карта, стиснена автоенкодером у латентний простір",
        "прогноз у латентному просторі дешевший, ніж генерувати повну карту",
    ),
]
cw = (12.13 - 2 * 0.3) / 3
for k, (name, color, what, why) in enumerate(models):
    x = 0.6 + k * (cw + 0.3)
    card(s, x, 3.85, cw, 2.6)
    pill(s, x + 0.2, 4.05, name, color, size=12)
    text(
        s,
        x + 0.2,
        4.55,
        cw - 0.4,
        2.0,
        [
            [("Що: ", {"bold": True, "color": MUTED}), (what, {})],
            [("Чому: ", {"bold": True, "color": MUTED}), (why, {})],
        ],
        size=16,
        space_after=10,
    )

# --------------------------------------------------------------------------- 3. metrics

s = new_slide(
    3,
    "ТЕХНІЧНА ЧАСТИНА · МЕТРИКИ Й БЕЙЗЛАЙНИ",
    "Як ми міряємо: метрики та бейзлайни",
    "Модель має сенс, лише якщо вона точніша за прості бейзлайни",
    notes=(
        "Доповідає: Олег. 40 с. Для траєкторії нашого авто міряємо L2 — відхилення в метрах через "
        "1, 2 і 3 секунди — і collision rate, частку зіткнень з іншими агентами. Для агентів — ADE "
        "і FDE, середню та кінцеву похибку прогнозу, і miss rate — частку промахів понад 2 метри. "
        "Усе реалізовано в коді з тестами. Бейзлайни: zero-motion (сцена не змінюється) і constant "
        "velocity (рух з останньою виміряною швидкістю). Ми стежили, щоб бейзлайни бачили лише "
        "минуле, а тест був на сценах, яких модель не бачила."
    ),
)
table(
    s,
    0.6,
    2.2,
    7.6,
    [
        ["Метрика", "Що вимірює", "Для чого"],
        ["L2 @ 1/2/3 с", "відхилення траєкторії авто, м", "BEV, усі"],
        ["Collision rate", "перетин авто з іншими агентами", "усі"],
        ["ADE / FDE / Miss", "похибка прогнозу агентів, м; промахи >2 м", "object-centric"],
        ["Chamfer distance", "збіг хмар точок", "BEV, latent BEV"],
    ],
    col_widths=[2.1, 3.7, 1.8],
    row_h=0.62,
    size=15,
    left_cols=(0, 1),
)
card(s, 8.6, 2.2, 4.13, 3.1)
text(s, 8.85, 2.4, 3.7, 0.4, [[("Бейзлайни", {"bold": True, "color": SKY})]], size=18)
text(
    s,
    8.85,
    2.95,
    3.7,
    2.3,
    [
        [("Zero-motion / Persistence: ", {"bold": True}), ("сцена не змінюється", {})],
        [("Constant velocity (CV): ", {"bold": True}), ("остання виміряна швидкість", {})],
        [("Ego-status MLP (план): ", {"bold": True}), ("чи не вирішує все сама швидкість его", {})],
    ],
    size=16,
    space_after=12,
    bullets=True,
)
card(s, 0.6, 5.55, 12.13, 1.15, fill=RGBColor(0x0F, 0x2A, 0x24), line=GREEN)
text(s, 0.85, 5.68, 3.0, 0.4, [[("Що ми врахували", {"bold": True, "color": GREEN})]], size=17)
text(
    s,
    0.85,
    6.1,
    11.7,
    0.5,
    [
        "бейзлайни бачать лише минуле  ·  тест — на 2 сценах, яких модель не бачила  ·  "
        "однакові кадри й метрики для всіх представлень"
    ],
    size=15,
)

# --------------------------------------------------------------------------- 4. BEV results

s = new_slide(
    4,
    "РЕЗУЛЬТАТИ · BEV",
    "BEV: на навчальних сценах точна, на нових — поки ні",
    "L2 у метрах на mini_val: 67 кадрів 2 сцен, яких модель не бачила (менше — краще)",
    badge_color=AMBER,
    notes=(
        "Доповідає: Владислав. 45 с. На навчальних сценах модель помиляється в середньому лише "
        "на ~0.1 м — у кілька разів менше, ніж constant velocity (≈1 м на тих самих сценах). Але на "
        "двох нових сценах її похибка 2.81 м, а в constant velocity 2.45 м. Тобто модель "
        "запам'ятала навчальні приклади, а не навчилася узагальнювати — це перенавчання. Причини: "
        "лише 267 навчальних прикладів, а енкодер камер і лідара ще не навчений. На анімації: "
        "жовте — як їхав водій, зелене — прогноз моделі; видно, що модель не завжди помічає, що "
        "водій гальмує."
    ),
)
table(
    s,
    0.6,
    2.2,
    5.6,
    [
        ["Метод", "1 с", "2 с", "3 с", "Сер."],
        ["Zero-motion", "5.23", "10.43", "15.55", "10.40"],
        ["Constant velocity", "0.69", "2.20", "4.46", "2.45"],
        ["BEVFusion + EgoPlanner", "0.80", "2.56", "5.06", "2.81"],
    ],
    col_widths=[2.4, 0.8, 0.8, 0.8, 0.8],
    row_h=0.55,
    size=15,
    highlight_row=2,
)
card(s, 0.6, 4.6, 5.6, 2.1)
text(s, 0.85, 4.75, 5.1, 0.4, [[("Що ми бачимо", {"bold": True, "color": AMBER})]], size=17)
text(
    s,
    0.85,
    5.2,
    5.1,
    1.45,
    [
        "Train: ~0.1 м — у кілька разів точніше за CV",
        "Val: 2.81 м проти 2.45 м у CV",
        "Перенавчання: мало даних, енкодер ще не навчений",
    ],
    size=15,
    bullets=True,
    space_after=5,
)
picture(s, FIGURES / "progress_bev_drive.gif", 6.5, 2.2, w=6.23)
text(
    s,
    6.5,
    2.2 + 6.23 * 385 / 684 + 0.12,
    6.23,
    0.6,
    ["Тестова сцена: жовте — водій, зелене — прогноз (у PPTX — анімація)"],
    size=13,
    color=MUTED,
)

# --------------------------------------------------------------------------- 5. OC results

s = new_slide(
    5,
    "РЕЗУЛЬТАТИ · OBJECT-CENTRIC",
    "Object-centric: сцена без втрат, CV — сильний бейзлайн",
    "Усі 10 сцен mini, 16\u00a0295 вікон агентів, похибка в метрах (менше — краще)",
    badge_color=GREEN,
    notes=(
        "Доповідає: Юрій. 45 с. Кожну сцену перетворюємо на тензор: кадри × агенти × 11 ознак; "
        "усі анотації потрапляють у тензор. Constant velocity у 4 рази точніша за persistence "
        "за FDE на 3 с — саме її має обійти наша навчена модель. На картинці пунктир — прогноз "
        "constant velocity, жовте — реальний рух: вони розходяться на поворотах і при гальмуванні."
    ),
)
card(s, 0.6, 2.2, 6.5, 1.0)
text(
    s,
    0.85,
    2.35,
    6.0,
    1.0,
    [
        [
            ("Тензор сцени: ", {"bold": True, "color": MUTED}),
            ("[кадри, агенти, 11 ознак], напр. [39, 227, 11]", {}),
        ],
        [
            ("Перевірка: ", {"bold": True, "color": MUTED}),
            ("усі 18\u00a0538 анотацій потрапили в тензор", {}),
        ],
    ],
    size=15,
    space_after=6,
)
table(
    s,
    0.6,
    3.5,
    6.5,
    [
        ["Бейзлайн", "ADE 1 с", "ADE 3 с", "FDE 3 с", "Miss 3 с"],
        ["Persistence", "1.07", "2.33", "3.97", "33.2%"],
        ["Constant velocity", "0.18", "0.50", "0.98", "10.7%"],
    ],
    col_widths=[2.1, 1.1, 1.1, 1.1, 1.1],
    row_h=0.6,
    size=16,
    highlight_row=2,
)
text(
    s,
    0.6,
    5.55,
    6.5,
    0.9,
    [
        [
            ("Висновок: ", {"bold": True, "color": GREEN}),
            ("CV у 4× точніша за persistence; помиляється на поворотах і гальмуванні", {}),
        ]
    ],
    size=16,
)
picture(s, FIGURES / "progress_object_forecasts.png", 7.55, 2.2, h=4.25)
text(
    s,
    7.55,
    6.5,
    5.2,
    0.4,
    ["Пунктир — прогноз constant velocity, жовте — реальний рух"],
    size=13,
    color=MUTED,
)

# --------------------------------------------------------------------------- 6. progress

s = new_slide(
    6,
    "ПРОГРЕС",
    "Що виконано з пропоузалу і хто що робив",
    "Два з трьох представлень уже дають результати; третє — наступний етап",
    notes=(
        "Доповідає: Олег. 45 с. З пропоузалу готові: підготовка даних, спільні метрики, "
        "бейзлайни і object-centric представлення. BEV-модель зібрана й навчена, але на нових "
        "сценах ще не краща за constant velocity. Latent BEV, стрес-тести і демо ще не почали: "
        "спершу хотіли, щоб порівняння було чесним і однаковим для всіх. Справа — хто що зробив."
    ),
)
items = [
    ("ГОТОВО", GREEN, "Дані nuScenes: keyframes, офіційні спліти"),
    ("ГОТОВО", GREEN, "Спільні метрики (L2, collision, Chamfer, FID, ADE/FDE) + тести"),
    ("ГОТОВО", GREEN, "Бейзлайни: zero-motion, persistence, constant velocity"),
    ("ГОТОВО", GREEN, "Object-centric представлення: без втрат, 10 сцен"),
    ("ВРАХОВАНО", SKY, "Чесне порівняння: лише минуле, тест на нових сценах"),
    ("ЧАСТКОВО", AMBER, "BEV-модель навчена, на val поки ≈ CV"),
    ("НЕ ПОЧАТО", ROSE, "Latent BEV, стрес-тести, демо — наступний етап"),
]
card(s, 0.6, 2.2, 7.0, 4.5)
for k, (label, color, desc) in enumerate(items):
    y = 2.4 + k * 0.6
    pill(s, 0.85, y + 0.1, label, color, w=1.5, size=11)
    text(s, 2.5, y, 4.95, 0.5, [desc], size=15, anchor=MSO_ANCHOR.MIDDLE, space_after=0)
people = [
    (
        "Олег Гомбош",
        "Репозиторій і CI, конфіг-схема, метрики L2 / collision / Chamfer / FID з тестами, "
        "persistence-бейзлайн",
    ),
    (
        "Владислав Васильченко",
        "nuScenes-лоадер, BEVFusion-енкодер і EgoPlanner, візуалізації (камера, BEV, GIF), "
        "слайди мідтерму",
    ),
    (
        "Юрій Бобас",
        "Object-centric представлення, ADE/FDE, ноутбук; аудит і виправлення методології BEV, "
        "навчання планера",
    ),
]
for k, (name, what) in enumerate(people):
    y = 2.2 + k * 1.55
    card(s, 7.9, y, 4.83, 1.4)
    text(
        s,
        8.15,
        y + 0.15,
        4.4,
        1.15,
        [[(name, {"bold": True, "color": SKY, "size": 17})], what],
        size=14,
        space_after=4,
    )

# --------------------------------------------------------------------------- 7. plan

s = new_slide(
    7,
    "ПОДАЛЬШИЙ ПЛАН",
    "План до фіналу: експерименти, план Б і що далі",
    "Для всіх трьох представлень — ті самі спліти, горизонти й бейзлайни",
    badge_color=VIOLET,
    notes=(
        "Доповідає: Юрій. 50 с. Спершу боремося з перенавчанням BEV: early stopping на окремій "
        "тренувальній сцені, LiDAR-растр і pretrained ResNet для камер. Потім навчена "
        "object-centric модель і спільна BEV-карта зайнятості, на якій порівнюємо всі "
        "представлення. Далі latent BEV — автоенкодер і прогноз у латентному просторі — та частина "
        "trainval. Наприкінці стрес-тести, демо і звіт. План Б: якщо дифузія не вміщається в GPU — "
        "простіший прогноз у латентному просторі; якщо енкодер не вчиться — будуємо BEV-карту "
        "прямо з анотацій і порівнюємо лише прогноз."
    ),
)
phases = [
    (
        "Листопад I",
        "BEV без перенавчання",
        "early stopping, LiDAR-растр (щільність, висота), pretrained ResNet",
    ),
    (
        "Листопад II",
        "Object-centric модель",
        "нейромережа по треках агентів; спільна BEV-карта зайнятості",
    ),
    (
        "Грудень I",
        "Latent BEV + масштаб",
        "автоенкодер + прогноз у латентному просторі; до 150 сцен trainval",
    ),
    ("Грудень II", "Стрес-тести і демо", "зсув его на 0.5 м, чутливість до дії водія, демо, звіт"),
]
pw = (12.13 - 3 * 0.25) / 4
for k, (when, head, body) in enumerate(phases):
    x = 0.6 + k * (pw + 0.25)
    card(s, x, 2.2, pw, 2.2)
    pill(s, x + 0.2, 2.38, when, VIOLET, w=1.45, size=11)
    text(
        s,
        x + 0.2,
        2.85,
        pw - 0.4,
        1.5,
        [[(head, {"bold": True, "size": 16})], body],
        size=15,
        space_after=4,
    )
for k, (head, color, lines) in enumerate(
    [
        (
            "Дизайн експериментів",
            SKY,
            [
                "train: mini_train, тест: mini_val",
                "early stopping на окремій train-сцені",
                "3 seeds, порівняння з CV",
            ],
        ),
        (
            "План Б",
            AMBER,
            [
                "дифузія не вміщається в GPU → простіший прогноз у латентному просторі",
                "енкодер не вчиться → BEV-карта з анотацій",
            ],
        ),
        (
            "Якщо все вийде",
            GREEN,
            ["повний trainval", "демо: як прогноз змінюється від дії водія"],
        ),
    ]
):
    w = (12.13 - 2 * 0.25) / 3
    x = 0.6 + k * (w + 0.25)
    card(s, x, 4.6, w, 2.1)
    text(s, x + 0.2, 4.75, w - 0.4, 0.4, [[(head, {"bold": True, "color": color})]], size=17)
    text(s, x + 0.2, 5.2, w - 0.4, 1.45, lines, size=15, space_after=4, bullets=True)

# --------------------------------------------------------------------------- save

prs.save(OUT_PPTX)
print(f"saved {OUT_PPTX}")
soffice = shutil.which("soffice") or shutil.which("libreoffice")
if soffice:
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(HERE), str(OUT_PPTX)],
        check=True,
    )
    print(f"saved {OUT_PPTX.with_suffix('.pdf')}")
else:
    print("LibreOffice not found: export the .pptx to PDF from PowerPoint or Google Slides")
